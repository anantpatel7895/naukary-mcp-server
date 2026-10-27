"""Parse a PDF or DOCX resume into a MasterProfile.

Heuristic and section-driven. Resumes are wildly inconsistent, so every stage degrades
rather than fails: an unrecognized section still lands in `raw_text`, which
`MasterProfile.searchable_text()` includes -- so keyword matching never silently loses
content just because the structural parse missed it.
"""

from __future__ import annotations

import re
from pathlib import Path

from naukri_mcp.resume.profile import (
    Bullet,
    Contact,
    Education,
    Experience,
    MasterProfile,
    Project,
    Skill,
)

# --- section headings ------------------------------------------------------------

SECTION_ALIASES: dict[str, str] = {
    "summary": "summary",
    "professional summary": "summary",
    "profile": "summary",
    "objective": "summary",
    "about": "summary",
    "about me": "summary",
    "technical skill": "skills",
    "technical skills": "skills",
    "skills": "skills",
    "skill": "skills",
    "core skills": "skills",
    "technical expertise": "skills",
    "competency": "skills",
    "competencies": "skills",
    "core competencies": "skills",
    "areas of expertise": "skills",
    "work history": "experience",
    "work experience": "experience",
    "experience": "experience",
    "professional experience": "experience",
    "employment": "experience",
    "employment history": "experience",
    "career history": "experience",
    "projects": "projects",
    "project": "projects",
    "personal projects": "projects",
    "key projects": "projects",
    "academic projects": "projects",
    "education": "education",
    "academics": "education",
    "academic qualification": "education",
    "educational qualification": "education",
    "certifications": "certifications",
    "certification": "certifications",
    "courses": "certifications",
    "licenses": "certifications",
    "research publications": "publications",
    "publications": "publications",
    "papers": "publications",
    "awards": "awards",
    "achievements": "awards",
}

BULLET_CHARS = "•·▪●◦‣-–—*>"
BULLET_RE = re.compile(rf"^\s*[{re.escape(BULLET_CHARS)}]\s+")

MONTH_YEAR = r"(?:\d{1,2}[-/]\d{4}|[A-Z][a-z]{2,8}\.?\s+\d{4}|\d{4})"
DATE_RANGE_RE = re.compile(
    rf"({MONTH_YEAR})\s*(?:[-–—]|to)\s*({MONTH_YEAR}|Current|Present|Till\s+Date|Ongoing)",
    re.I,
)

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\d[\s-]?){9,12}\d")
URL_RE = re.compile(r"https?://\S+|(?:www\.)\S+")

# A skills line like "Programming Language: Python, Java"
SKILL_CATEGORY_RE = re.compile(r"^\s*([A-Za-z][A-Za-z /&+.\-]{2,45}?)\s*:\s*(.+)$")
SPLIT_SKILLS_RE = re.compile(r"[,;|/]|\s+and\s+")


def _slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s or "item"


# --- text extraction -------------------------------------------------------------


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix in {".docx", ".doc"}:
        if suffix == ".doc":
            raise ValueError(
                "Legacy .doc is not supported. Save as .docx or PDF and re-import."
            )
        return _extract_docx(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported resume format '{suffix}'. Use PDF, DOCX, TXT, or MD.")


def _extract_pdf(path: Path) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    # Pull table cells too -- plenty of resumes lay out skills in a table.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    lines.append(cell_text)
    return "\n".join(lines)


# --- sectioning ------------------------------------------------------------------


def reflow_lines(lines: list[str]) -> list[str]:
    """Rejoin lines that a PDF extractor hard-wrapped mid-sentence.

    Without this, one wrapped bullet becomes two bullets and one wrapped paragraph
    becomes several bogus "projects". The signal is line width: in a wrapped block every
    line but the last runs to nearly the full text width, so a near-full-width line is
    almost certainly continued on the next one. The width is measured from the document
    itself rather than hardcoded, so it adapts to margins, font size, and page geometry.

    Only ever call this on PDF-extracted text. DOCX paragraphs are already logical units,
    and reflowing them merges an entire section into one line.
    """
    widths = sorted(len(ln.strip()) for ln in lines if ln.strip())
    if not widths:
        return lines
    # 90th percentile approximates the text-column width while ignoring outliers.
    wrap_width = widths[min(int(len(widths) * 0.9), len(widths) - 1)]
    threshold = wrap_width * 0.80
    if threshold < 25:
        # Very short lines throughout: there is no wrapping to undo, and guessing would
        # corrupt the structure.
        return lines

    # Require positive evidence of wrapping. Genuine wrapped paragraphs leave a tight
    # cluster of lines just under the column width; a document with only a handful of
    # long lines among many short ones is laid out deliberately, not wrapped. Absolute
    # count catches full documents; the proportion catches short excerpts.
    cluster = sum(1 for w in widths if w >= wrap_width * 0.92)
    if cluster < 4 and cluster / len(widths) < 0.3:
        return lines

    out: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            out.append("")
            continue

        starts_new_block = bool(
            BULLET_RE.match(raw)
            or _normalize_heading(line)
            or DATE_RANGE_RE.search(line)
            or SKILL_CATEGORY_RE.match(line)
        )
        if not starts_new_block and out and out[-1].strip():
            prev = out[-1].strip()
            prev_is_wrapped = len(prev) >= threshold and not _normalize_heading(prev)
            if prev_is_wrapped:
                out[-1] = f"{prev} {line}"
                continue
        out.append(line)
    return out


def _normalize_heading(line: str) -> str | None:
    """Return the canonical section name if `line` looks like a heading."""
    cleaned = line.strip().strip(":").strip()
    if not cleaned or len(cleaned) > 45:
        return None
    if BULLET_RE.match(line):
        return None
    key = re.sub(r"\s+", " ", cleaned.lower())
    return SECTION_ALIASES.get(key)


def split_sections(lines: list[str]) -> tuple[list[str], dict[str, list[str]]]:
    """Split into (header_lines_before_first_section, {section: lines})."""
    header: list[str] = []
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for line in lines:
        name = _normalize_heading(line)
        if name:
            current = name
            sections.setdefault(current, [])
            continue
        if current is None:
            header.append(line)
        else:
            sections[current].append(line)
    return header, sections


# --- per-section parsers ---------------------------------------------------------


def parse_contact(header_lines: list[str], full_text: str) -> Contact:
    contact = Contact()
    nonempty = [ln.strip() for ln in header_lines if ln.strip()]
    if nonempty:
        contact.name = nonempty[0][:120]

    # Search the header first, then the whole document as a fallback.
    for scope in ("\n".join(nonempty), full_text):
        if not contact.email and (m := EMAIL_RE.search(scope)):
            contact.email = m.group(0)
        if not contact.phone:
            for cand in PHONE_RE.findall(scope):
                digits = re.sub(r"\D", "", cand)
                # Avoid matching year ranges / CGPA / pin codes.
                if 10 <= len(digits) <= 13:
                    contact.phone = cand.strip()
                    break
        if contact.email and contact.phone:
            break

    for url in URL_RE.findall("\n".join(nonempty)) or URL_RE.findall(full_text):
        low = url.lower()
        if "linkedin" in low and not contact.linkedin:
            contact.linkedin = url
        elif "github" in low and not contact.github:
            contact.github = url
        elif not contact.website and "linkedin" not in low and "github" not in low:
            contact.website = url

    # The line after the name is usually a headline containing a location.
    if len(nonempty) > 1:
        headline = nonempty[1]
        if "," in headline:
            contact.location = headline.rsplit(",", 1)[-1].strip()[:120]
    return contact


def parse_skills(lines: list[str]) -> list[Skill]:
    skills: list[Skill] = []
    seen: set[str] = set()

    def add(name: str, category: str) -> None:
        name = name.strip(" .;:•-")
        # Drop empties, single characters, and runaway sentence fragments.
        if not name or len(name) < 2 or len(name) > 60:
            return
        key = name.lower()
        if key in seen:
            return
        seen.add(key)
        skills.append(Skill(ref=f"skill.{_slug(name)}", name=name, category=category))

    for raw in lines:
        line = BULLET_RE.sub("", raw).strip()
        if not line:
            continue
        if m := SKILL_CATEGORY_RE.match(line):
            category, payload = m.group(1).strip(), m.group(2)
        else:
            category, payload = "", line
        for token in SPLIT_SKILLS_RE.split(payload):
            add(token, category)
    return skills


def parse_experience(lines: list[str]) -> list[Experience]:
    experiences: list[Experience] = []
    idx = 0
    pending_bullets: list[str] = []

    def flush_bullets() -> None:
        nonlocal pending_bullets
        if experiences and pending_bullets:
            exp = experiences[-1]
            for text in pending_bullets:
                exp.bullets.append(
                    Bullet(ref=f"{exp.ref}.bullet.{len(exp.bullets)}", text=text)
                )
        pending_bullets = []

    for raw in lines:
        line = raw.strip()
        if not line:
            continue

        if BULLET_RE.match(raw):
            pending_bullets.append(BULLET_RE.sub("", raw).strip())
            continue

        date_match = DATE_RANGE_RE.search(line)
        if date_match:
            # New role header. Title is whatever precedes the date/separator.
            flush_bullets()
            title = line[: date_match.start()].strip(" |:–-—,")
            exp = Experience(
                ref=f"exp.{idx}",
                company="",
                title=title[:200] or "Role",
                start=date_match.group(1).strip(),
                end=date_match.group(2).strip(),
            )
            experiences.append(exp)
            idx += 1
            continue

        if experiences and not experiences[-1].company and not pending_bullets:
            # The line right after the header is the company (and often location).
            exp = experiences[-1]
            if "," in line:
                company, _, location = line.rpartition(",")
                exp.company = company.strip()[:200]
                exp.location = location.strip()[:120]
            else:
                exp.company = line[:200]
            continue

        # Any other prose inside the section: treat as an unmarked bullet.
        if experiences:
            pending_bullets.append(line)

    flush_bullets()
    return experiences


def parse_projects(lines: list[str]) -> list[Project]:
    projects: list[Project] = []
    idx = 0
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if BULLET_RE.match(raw):
            text = BULLET_RE.sub("", raw).strip()
            if projects:
                proj = projects[-1]
                proj.bullets.append(
                    Bullet(ref=f"{proj.ref}.bullet.{len(proj.bullets)}", text=text)
                )
            continue

        # After reflow, each paragraph is a whole unit: a title is short and lacks
        # terminal sentence punctuation, whereas a description is a full sentence.
        is_title = len(line) < 120 and not line.rstrip('"').endswith((".", ";", ","))
        if is_title or not projects:
            projects.append(Project(ref=f"proj.{idx}", name=line[:250]))
            idx += 1
        else:
            proj = projects[-1]
            proj.description = f"{proj.description} {line}".strip()
    return projects


def parse_education(lines: list[str]) -> list[Education]:
    entries: list[Education] = []
    idx = 0
    for raw in lines:
        line = BULLET_RE.sub("", raw).strip()
        if not line:
            continue

        date_match = DATE_RANGE_RE.search(line)
        looks_like_degree = bool(
            re.search(
                r"\b(bachelor|master|b\.?tech|m\.?tech|b\.?e\b|m\.?e\b|b\.?sc|m\.?sc|"
                r"mba|phd|diploma|\d{1,2}th|hsc|ssc|intermediate)\b",
                line,
                re.I,
            )
        )
        if date_match or looks_like_degree or not entries:
            degree, _, field = line.partition(":")
            entries.append(
                Education(
                    ref=f"edu.{idx}",
                    institution="",
                    degree=degree.split("|")[0].strip()[:200],
                    field=field.split("|")[0].strip()[:200],
                    start=date_match.group(1).strip() if date_match else "",
                    end=date_match.group(2).strip() if date_match else "",
                    details=line[:400],
                )
            )
            idx += 1
        elif entries and not entries[-1].institution:
            entries[-1].institution = line[:250]
    return entries


# --- entry point -----------------------------------------------------------------


def parse_resume(path: str | Path) -> MasterProfile:
    resume_path = Path(path).expanduser()
    if not resume_path.exists():
        raise FileNotFoundError(f"No resume at {resume_path}")

    text = extract_text(resume_path)
    if not text.strip():
        raise ValueError(
            f"Extracted no text from {resume_path.name}. If it is a scanned image, "
            "it needs OCR before import."
        )

    # Only PDFs are hard-wrapped by extraction; DOCX/TXT paragraphs are already units.
    lines = text.splitlines()
    if resume_path.suffix.lower() == ".pdf":
        lines = reflow_lines(lines)
    header, sections = split_sections(lines)

    summary_lines = [ln.strip() for ln in sections.get("summary", []) if ln.strip()]
    profile = MasterProfile(
        contact=parse_contact(header, text),
        summary=" ".join(summary_lines)[:2000],
        skills=parse_skills(sections.get("skills", [])),
        experience=parse_experience(sections.get("experience", [])),
        projects=parse_projects(sections.get("projects", [])),
        education=parse_education(sections.get("education", [])),
        certifications=[
            BULLET_RE.sub("", ln).strip()
            for ln in sections.get("certifications", [])
            if ln.strip()
        ],
        raw_text=text,
        source_file=str(resume_path),
    )
    _link_skill_evidence(profile)
    return profile


def _link_skill_evidence(profile: MasterProfile) -> None:
    """Record which bullets demonstrate each skill.

    A skill backed by a bullet is much stronger evidence than one that only appears in
    a comma-separated list, and the gap report surfaces that difference.
    """
    bullets: list[tuple[str, str]] = []
    for exp in profile.experience:
        bullets.extend((b.ref, b.text.lower()) for b in exp.bullets)
    for proj in profile.projects:
        bullets.extend((b.ref, b.text.lower()) for b in proj.bullets)
        if proj.description:
            bullets.append((proj.ref, f"{proj.name} {proj.description}".lower()))

    for skill in profile.skills:
        needle = skill.name.lower()
        if len(needle) < 3:
            continue
        pattern = re.compile(rf"(?<![a-z0-9]){re.escape(needle)}(?![a-z0-9])")
        skill.evidenced_by = [ref for ref, text in bullets if pattern.search(text)]
