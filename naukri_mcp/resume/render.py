"""Validate proposed tailoring, then render an ATS-safe resume.

`validate_tailoring` is the safety boundary of this whole server. It refuses to render
any bullet that is not traceable to the MasterProfile, or that asserts a keyword the gap
report classified as `missing`. The worst outcome is therefore a resume that scores
lower, never one that claims something untrue.
"""

from __future__ import annotations

import re
from pathlib import Path

from pydantic import BaseModel, Field

from naukri_mcp.resume.keywords import Bucket, GapReport, _PATTERNS, find_skills
from naukri_mcp.resume.profile import MasterProfile

# ATS parsers are most reliable with one standard serif/sans face and no exotic glyphs.
ATS_FONT = "Calibri"
ATS_FONT_SIZE = 10.5
ATS_NAME_SIZE = 18


class TailoredBullet(BaseModel):
    """A rewritten achievement line, with its provenance."""

    #: Which MasterProfile entry this is a rewrite of. Must exist in the evidence index.
    evidence_ref: str
    text: str
    #: Optional: the JD keywords this rewrite is meant to surface.
    surfaces: list[str] = Field(default_factory=list)


class TailoringProposal(BaseModel):
    """Everything the host LLM proposes changing. Nothing else is editable."""

    summary: str = ""
    #: Rewritten experience bullets, replacing the originals they cite.
    bullets: list[TailoredBullet] = Field(default_factory=list)
    #: Reordering only -- these must all already exist in the profile.
    skill_order: list[str] = Field(default_factory=list)
    label: str = ""


class ValidationError(BaseModel):
    where: str
    problem: str
    offending_text: str = ""


class ValidationResult(BaseModel):
    ok: bool
    errors: list[ValidationError] = Field(default_factory=list)
    #: Keywords the proposal legitimately surfaces, for A/B analytics later.
    surfaced_keywords: list[str] = Field(default_factory=list)

    def raise_for_status(self) -> None:
        if not self.ok:
            lines = [
                f"  - [{e.where}] {e.problem}" + (f"\n      {e.offending_text!r}" if e.offending_text else "")
                for e in self.errors
            ]
            raise TailoringRejected(
                "Refusing to render this tailoring because it is not supported by your "
                "resume:\n" + "\n".join(lines)
            )


class TailoringRejected(ValueError):
    """The proposal asserted something the MasterProfile does not support."""


def _asserts_keyword(text: str, keyword: str) -> bool:
    return any(p.search(text) for p in _PATTERNS.get(keyword, []))


def validate_tailoring(
    proposal: TailoringProposal,
    profile: MasterProfile,
    report: GapReport | None = None,
) -> ValidationResult:
    """Check every proposed edit against the profile.

    Three rules, in order of importance:
      1. No bullet may assert a keyword bucketed `missing` -- that is fabrication.
      2. Every bullet must cite an `evidence_ref` that exists in the profile.
      3. Skill reordering may not introduce a skill that is not already present.
    """
    errors: list[ValidationError] = []
    index = profile.evidence_index()
    forbidden = report.forbidden_keywords if report else set()
    # `adjacent` keywords may be alluded to, but never asserted as owned skills.
    adjacent_only = (
        {f.keyword for f in report.by_bucket(Bucket.adjacent)} if report else set()
    )

    surfaced: set[str] = set()

    for i, bullet in enumerate(proposal.bullets):
        where = f"bullets[{i}]"
        text = bullet.text.strip()

        if not text:
            errors.append(ValidationError(where=where, problem="Bullet text is empty."))
            continue

        if not bullet.evidence_ref:
            errors.append(
                ValidationError(
                    where=where,
                    problem="No evidence_ref. Every bullet must cite the resume entry it rewrites.",
                    offending_text=text,
                )
            )
        elif bullet.evidence_ref not in index:
            errors.append(
                ValidationError(
                    where=where,
                    problem=(
                        f"evidence_ref {bullet.evidence_ref!r} does not exist in your resume. "
                        "Use a ref from analyze_job_match / resume://master."
                    ),
                    offending_text=text,
                )
            )

        for keyword in sorted(forbidden):
            if _asserts_keyword(text, keyword):
                errors.append(
                    ValidationError(
                        where=where,
                        problem=(
                            f"Asserts {keyword!r}, which is absent from your resume "
                            "(bucket=missing). This would be fabrication."
                        ),
                        offending_text=text,
                    )
                )

        for keyword in sorted(adjacent_only):
            if _asserts_keyword(text, keyword):
                errors.append(
                    ValidationError(
                        where=where,
                        problem=(
                            f"Asserts {keyword!r} as owned experience, but you only have "
                            "adjacent skills. Reword to name what you actually used."
                        ),
                        offending_text=text,
                    )
                )

        surfaced.update(find_skills(text))

    if proposal.summary:
        for keyword in sorted(forbidden | adjacent_only):
            if _asserts_keyword(proposal.summary, keyword):
                errors.append(
                    ValidationError(
                        where="summary",
                        problem=f"Summary asserts {keyword!r}, which your resume does not support.",
                        offending_text=proposal.summary[:200],
                    )
                )
        surfaced.update(find_skills(proposal.summary))

    known_skills = {s.name.lower() for s in profile.skills}
    for name in proposal.skill_order:
        if name.lower() not in known_skills:
            errors.append(
                ValidationError(
                    where="skill_order",
                    problem=(
                        f"{name!r} is not in your resume's skills. Reordering may not add "
                        "new skills."
                    ),
                    offending_text=name,
                )
            )

    have = find_skills(profile.searchable_text())
    return ValidationResult(
        ok=not errors,
        errors=errors,
        surfaced_keywords=sorted(surfaced & have),
    )


def apply_proposal(profile: MasterProfile, proposal: TailoringProposal) -> MasterProfile:
    """Return a copy of `profile` with the proposal's edits applied.

    Assumes validation already passed. Bullets replace the entry they cite; skills are
    reordered; nothing is invented.
    """
    tailored = profile.model_copy(deep=True)

    if proposal.summary:
        tailored.summary = proposal.summary

    replacements = {b.evidence_ref: b.text for b in proposal.bullets}
    for exp in tailored.experience:
        for bullet in exp.bullets:
            if bullet.ref in replacements:
                bullet.text = replacements[bullet.ref]
    for proj in tailored.projects:
        for bullet in proj.bullets:
            if bullet.ref in replacements:
                bullet.text = replacements[bullet.ref]

    if proposal.skill_order:
        rank = {name.lower(): i for i, name in enumerate(proposal.skill_order)}
        tailored.skills.sort(key=lambda s: rank.get(s.name.lower(), len(rank)))

    return tailored


# --- DOCX rendering --------------------------------------------------------------


def render_docx(profile: MasterProfile, out_path: Path) -> Path:
    """Write a single-column, ATS-parseable DOCX.

    Deliberately plain: no tables, text boxes, headers/footers, columns, or images. Those
    are the constructs that make ATS parsers drop or scramble content.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()

    style = document.styles["Normal"]
    style.font.name = ATS_FONT
    style.font.size = Pt(ATS_FONT_SIZE)
    style.paragraph_format.space_after = Pt(2)

    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(36)
        section.left_margin = section.right_margin = Pt(40)

    def heading(text: str) -> None:
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(ATS_FONT_SIZE + 1)

    def body(text: str, *, bold: bool = False, bullet: bool = False) -> None:
        p = document.add_paragraph(style="List Bullet" if bullet else None)
        run = p.add_run(text)
        run.bold = bold

    contact = profile.contact
    name_p = document.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(contact.name or "Candidate")
    name_run.bold = True
    name_run.font.size = Pt(ATS_NAME_SIZE)

    # Contact details go in the body, never a header/footer -- many ATS never read those.
    details = [contact.phone, contact.email, contact.location,
               contact.linkedin, contact.github, contact.website]
    detail_p = document.add_paragraph()
    detail_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail_p.add_run(" | ".join(d for d in details if d))

    if profile.summary:
        heading("Summary")
        body(profile.summary)

    if profile.skills:
        heading("Skills")
        grouped: dict[str, list[str]] = {}
        for skill in profile.skills:
            grouped.setdefault(skill.category or "Other", []).append(skill.name)
        for category, names in grouped.items():
            p = document.add_paragraph()
            label = p.add_run(f"{category}: ")
            label.bold = True
            p.add_run(", ".join(names))

    if profile.experience:
        heading("Experience")
        for exp in profile.experience:
            # Title and dates share one line, company on the next. Resume parsers anchor
            # a role on the line carrying the date range, so splitting them makes the
            # company line look like the job title -- tests/test_render_roundtrip.py
            # caught exactly that.
            dates = " - ".join(d for d in (exp.start, exp.end) if d)
            body(" | ".join(x for x in (exp.title, dates) if x), bold=True)
            meta = ", ".join(m for m in (exp.company, exp.location) if m)
            if meta:
                body(meta)
            for bullet in exp.bullets:
                body(bullet.text, bullet=True)

    if profile.projects:
        heading("Projects")
        for proj in profile.projects:
            body(proj.name, bold=True)
            if proj.tech:
                body("Tech: " + ", ".join(proj.tech))
            if proj.description:
                body(proj.description)
            for bullet in proj.bullets:
                body(bullet.text, bullet=True)

    if profile.education:
        heading("Education")
        for edu in profile.education:
            dates = " - ".join(d for d in (edu.start, edu.end) if d)
            headline = " | ".join(
                x for x in (f"{edu.degree}: {edu.field}".strip(": "), dates) if x
            )
            body(headline, bold=True)
            if edu.institution:
                body(edu.institution)

    if profile.certifications:
        heading("Certifications")
        for cert in profile.certifications:
            body(cert, bullet=True)

    document.save(str(out_path))
    return out_path


def render_pdf(profile: MasterProfile, out_path: Path) -> Path:
    """Plain-text-flow PDF for human readers. DOCX stays the ATS submission format."""
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    out_path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "b", parent=styles["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.5
    )
    name_style = ParagraphStyle(
        "n", parent=base, fontName="Helvetica-Bold", fontSize=17, leading=20, alignment=TA_CENTER
    )
    center = ParagraphStyle("c", parent=base, alignment=TA_CENTER)
    head = ParagraphStyle(
        "h", parent=base, fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=8, spaceAfter=3
    )
    bullet_style = ParagraphStyle("l", parent=base, leftIndent=10, bulletIndent=2)

    def esc(text: str) -> str:
        return (
            text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )

    flow: list[object] = []
    contact = profile.contact
    flow.append(Paragraph(esc(contact.name or "Candidate"), name_style))
    details = [contact.phone, contact.email, contact.location,
               contact.linkedin, contact.github, contact.website]
    flow.append(Paragraph(esc(" | ".join(d for d in details if d)), center))
    flow.append(Spacer(1, 4))

    def section(title: str) -> None:
        flow.append(Paragraph(esc(title.upper()), head))

    if profile.summary:
        section("Summary")
        flow.append(Paragraph(esc(profile.summary), base))

    if profile.skills:
        section("Skills")
        grouped: dict[str, list[str]] = {}
        for skill in profile.skills:
            grouped.setdefault(skill.category or "Other", []).append(skill.name)
        for category, names in grouped.items():
            flow.append(
                Paragraph(f"<b>{esc(category)}:</b> {esc(', '.join(names))}", base)
            )

    if profile.experience:
        section("Experience")
        for exp in profile.experience:
            # Same title|dates-then-company ordering as the DOCX, for the same reason.
            dates = " - ".join(d for d in (exp.start, exp.end) if d)
            headline = " | ".join(x for x in (exp.title, dates) if x)
            flow.append(Paragraph(f"<b>{esc(headline)}</b>", base))
            meta = ", ".join(m for m in (exp.company, exp.location) if m)
            if meta:
                flow.append(Paragraph(esc(meta), base))
            for bullet in exp.bullets:
                flow.append(Paragraph(esc(bullet.text), bullet_style, bulletText="•"))

    if profile.projects:
        section("Projects")
        for proj in profile.projects:
            flow.append(Paragraph(f"<b>{esc(proj.name)}</b>", base))
            if proj.tech:
                flow.append(Paragraph(esc("Tech: " + ", ".join(proj.tech)), base))
            if proj.description:
                flow.append(Paragraph(esc(proj.description), base))
            for bullet in proj.bullets:
                flow.append(Paragraph(esc(bullet.text), bullet_style, bulletText="•"))

    if profile.education:
        section("Education")
        for edu in profile.education:
            dates = " - ".join(d for d in (edu.start, edu.end) if d)
            headline = " | ".join(
                x for x in (f"{edu.degree}: {edu.field}".strip(": "), dates) if x
            )
            flow.append(Paragraph(f"<b>{esc(headline)}</b>", base))
            if edu.institution:
                flow.append(Paragraph(esc(edu.institution), base))

    if profile.certifications:
        section("Certifications")
        for cert in profile.certifications:
            flow.append(Paragraph(esc(cert), bullet_style, bulletText="•"))

    SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        topMargin=13 * mm,
        bottomMargin=13 * mm,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        title=f"{contact.name} Resume".strip(),
    ).build(flow)
    return out_path


_SAFE_NAME_RE = re.compile(r"[^A-Za-z0-9._-]+")


def safe_filename(*parts: str, suffix: str) -> str:
    stem = "_".join(p for p in parts if p) or "resume"
    stem = _SAFE_NAME_RE.sub("-", stem).strip("-_.")[:90] or "resume"
    return f"{stem}{suffix}"
