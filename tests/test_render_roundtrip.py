"""ATS-safety checks on the rendered resume.

The round-trip test is the proxy for ATS parseability: if our own section-driven parser
can recover the structure from the generated DOCX, a commercial parser very likely can
too. The structural assertions pin the layout rules that actually break ATS parsers.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from naukri_mcp.resume.parse import parse_resume
from naukri_mcp.resume.profile import (
    Bullet,
    Contact,
    Education,
    Experience,
    MasterProfile,
    Project,
    Skill,
)
from naukri_mcp.resume.render import render_docx, render_pdf, safe_filename


@pytest.fixture
def profile() -> MasterProfile:
    return MasterProfile(
        contact=Contact(
            name="Priya Sharma",
            email="priya@example.com",
            phone="9876543210",
            location="Pune",
            linkedin="https://linkedin.com/in/priya",
        ),
        summary="Backend engineer building Python services.",
        skills=[
            Skill(ref="skill.python", name="Python", category="Programming Language"),
            Skill(ref="skill.fastapi", name="FastAPI", category="Web Development Frameworks"),
            Skill(ref="skill.mysql", name="MySQL", category="Database"),
        ],
        experience=[
            Experience(
                ref="exp.0",
                company="Example Corp",
                title="Backend Engineer",
                start="01-2023",
                end="Current",
                location="Pune",
                bullets=[
                    Bullet(ref="exp.0.bullet.0", text="Built Python REST services with FastAPI."),
                    Bullet(ref="exp.0.bullet.1", text="Tuned MySQL queries for reporting."),
                ],
            )
        ],
        projects=[
            Project(
                ref="proj.0",
                name="Doc Search",
                description="Semantic search over internal documents.",
                tech=["Python", "FAISS"],
            )
        ],
        education=[
            Education(
                ref="edu.0",
                institution="Pune Institute of Technology",
                degree="Bachelor of Engineering",
                field="Computer Science",
                start="2017",
                end="2021",
            )
        ],
        certifications=["AWS Certified Cloud Practitioner"],
    )


def test_docx_roundtrips_through_our_own_parser(profile, tmp_path: Path):
    out = render_docx(profile, tmp_path / "r.docx")
    assert out.exists() and out.stat().st_size > 0

    reparsed = parse_resume(out)

    assert reparsed.contact.name == "Priya Sharma"
    assert reparsed.contact.email == "priya@example.com"
    assert reparsed.summary.startswith("Backend engineer")

    skill_names = {s.name.lower() for s in reparsed.skills}
    assert {"python", "fastapi", "mysql"} <= skill_names

    assert len(reparsed.experience) == 1
    exp = reparsed.experience[0]
    assert "Backend Engineer" in exp.title
    assert "Example Corp" in exp.company
    assert len(exp.bullets) == 2
    assert any("FastAPI" in b.text for b in exp.bullets)

    assert any("Doc Search" in p.name for p in reparsed.projects)
    assert reparsed.education


def test_docx_avoids_ats_hostile_constructs(profile, tmp_path: Path):
    """Tables, text boxes, and header/footer content are what ATS parsers drop."""
    import docx

    out = render_docx(profile, tmp_path / "r.docx")
    document = docx.Document(str(out))

    assert document.tables == [], "Tables scramble ATS column ordering."

    body_xml = document.element.body.xml
    assert "<w:txbxContent" not in body_xml, "Text boxes are invisible to many ATS."
    assert "<w:drawing" not in body_xml, "Images carry no parseable text."
    assert "w:cols w:num" not in body_xml, "Multi-column layout breaks reading order."

    # Contact details must live in the body, not a header/footer that ATS never reads.
    for section in document.sections:
        header_text = " ".join(p.text for p in section.header.paragraphs).strip()
        assert profile.contact.email not in header_text
    body_text = "\n".join(p.text for p in document.paragraphs)
    assert profile.contact.email in body_text
    assert profile.contact.phone in body_text


def test_docx_uses_standard_section_headings(profile, tmp_path: Path):
    """ATS parsers key off conventional heading words."""
    import docx

    out = render_docx(profile, tmp_path / "r.docx")
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs).upper()
    for heading in ("SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION"):
        assert heading in text


def test_pdf_renders(profile, tmp_path: Path):
    out = render_pdf(profile, tmp_path / "r.pdf")
    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")


def test_pdf_text_is_extractable(profile, tmp_path: Path):
    """A PDF whose text cannot be extracted is useless to an ATS."""
    import pdfplumber

    out = render_pdf(profile, tmp_path / "r.pdf")
    with pdfplumber.open(str(out)) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert "Priya Sharma" in text
    assert "FastAPI" in text


@pytest.mark.parametrize(
    "parts,expected_absent",
    [(("Priya Sharma", "Acme/Corp", "Sr. Engineer"), "/"), (("a" * 200,), None)],
)
def test_safe_filename_sanitizes(parts, expected_absent):
    name = safe_filename(*parts, suffix=".docx")
    assert name.endswith(".docx")
    assert len(name) <= 96
    if expected_absent:
        assert expected_absent not in name.replace(".docx", "")
